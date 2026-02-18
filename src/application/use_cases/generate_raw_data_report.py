import duckdb
import io
import datetime
from typing import Dict, Any, List
import pandas as pd
import openpyxl

class GenerateRawDataReport:
    def __init__(self, db_path: str = "data/metrics.duckdb"):
        self.db_path = db_path

    def _get_filter_summary(self, filters: Dict[str, Any]) -> Dict[str, str]:
        summary = {}
        start, end = filters.get('start_date'), filters.get('end_date')
        if start and end: summary['Rango de Fechas'] = f"{start} al {end}"
        elif start: summary['Rango de Fechas'] = f"Desde {start}"
        elif end: summary['Rango de Fechas'] = f"Hasta {end}"
        else: summary['Rango de Fechas'] = "Todo el periodo"
        return summary

    def generate_excel(self, filters: Dict[str, Any]) -> io.BytesIO:
        conn = duckdb.connect(self.db_path, read_only=True)
        try:
            # Base query with Joins
            # f.*, region_orig.name, region_dest.name, file.file_name
            query = """
                SELECT 
                    f.*,
                    r_orig.name as region_origen,
                    r_dest.name as region_destino,
                    fc.file_name as nombre_archivo
                FROM flights f
                LEFT JOIN region_airports ra_orig ON f.origen = ra_orig.icao_code
                LEFT JOIN regions r_orig ON ra_orig.region_id = r_orig.id
                LEFT JOIN region_airports ra_dest ON f.destino = ra_dest.icao_code
                LEFT JOIN regions r_dest ON ra_dest.region_id = r_dest.id
                LEFT JOIN file_processing_control fc ON f.file_id = fc.id
                WHERE 1=1
            """
            params = []
            
            # Apply Filters
            if filters.get('start_date'):
                query += " AND f.fecha >= ?"
                params.append(filters['start_date'])
            if filters.get('end_date'):
                query += " AND f.fecha <= ?"
                params.append(filters['end_date'])
            
            # List filters
            # Helper to add WHERE IN clauses
            def add_list_filter(key, col):
                items = filters.get(key, [])
                if items:
                    values = []
                    for x in items:
                        if isinstance(x, dict):
                             # Handle commonly used frontend headers
                            if 'value' in x: 
                                 val = x['value']
                                 if isinstance(val, dict) and 'icao_code' in val: values.append(str(val['icao_code']))
                                 else: values.append(str(val))
                            elif 'id' in x: values.append(str(x['id']))
                            elif 'label' in x: values.append(str(x['label'])) # fallback
                        else:
                            values.append(str(x))
                    
                    if values:
                        placeholders = ','.join(['?'] * len(values))
                        nonlocal query
                        query += f" AND {col} IN ({placeholders})"
                        params.extend(values)

            add_list_filter('origins', 'f.origen')
            add_list_filter('destinations', 'f.destino')
            add_list_filter('empresa', 'f.empresa')
            add_list_filter('tipo_aeronave', 'f.tipo_aeronave')
            add_list_filter('tipo_vuelo', 'f.tipo_vuelo')
            add_list_filter('callsign', 'f.callsign')
            add_list_filter('matriculas', 'f.matricula')
            
            # Additional logic for min/max level if needed, but usually those are not primary filters for this report?
            # Creating DF directly via DuckDB is faster
            # But we need parameters. 
            # DuckDB fetchdf is efficient.
            
            df = conn.execute(query, params).fetchdf()
            
            output = io.BytesIO()
            writer = pd.ExcelWriter(output, engine='openpyxl')
            
            if not df.empty:
                df.to_excel(writer, sheet_name='Data Cruda', index=False)
            else:
                 pd.DataFrame({'Info': ['No data found matching filters']}).to_excel(writer, sheet_name='Data Cruda', index=False)

            # Metadata Sheet
            summary = self._get_filter_summary(filters)
            if summary:
                ws_meta = writer.book.create_sheet('Filtros')
                ws_meta.append(['Filtro', 'Valor'])
                for k, v in summary.items():
                    ws_meta.append([k, v])
            
            writer.close()
            output.seek(0)
            return output

        except Exception as e:
            print(f"Error generating raw report: {e}")
            raise e
        finally:
            conn.close()
